"""
文档加载器

支持加载 PDF / Markdown / HTML 等格式文档，并转换为纯文本供后续处理。
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Iterator

from loguru import logger


# 支持的文件扩展名
SUPPORTED_EXTENSIONS = {".pdf", ".md", ".markdown", ".html", ".htm", ".txt", ".docx"}


def read_pdf(file_path: Path) -> str:
    """读取 PDF 文件"""
    from pypdf import PdfReader

    reader = PdfReader(str(file_path))
    pages = []
    for i, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
            pages.append(text)
        except Exception as e:
            logger.warning(f"PDF 第 {i + 1} 页读取失败: {e}")
    return "\n\n".join(pages)


def read_markdown(file_path: Path) -> str:
    """读取 Markdown 文件（保持原文，后续可转 HTML）"""
    return file_path.read_text(encoding="utf-8", errors="ignore")


def read_html(file_path: Path) -> str:
    """读取 HTML 文件并提取纯文本"""
    from bs4 import BeautifulSoup

    html = file_path.read_text(encoding="utf-8", errors="ignore")
    soup = BeautifulSoup(html, "html.parser")
    # 移除 script / style
    for tag in soup.find_all(["script", "style"]):
        tag.decompose()
    return soup.get_text(separator="\n", strip=True)


def read_txt(file_path: Path) -> str:
    """读取纯文本文件"""
    return file_path.read_text(encoding="utf-8", errors="ignore")


def read_docx(file_path: Path) -> str:
    """读取 Word 文件"""
    from docx import Document as DocxDocument

    doc = DocxDocument(str(file_path))
    return "\n".join(p.text for p in doc.paragraphs if p.text)


def load_document(file_path: str | Path) -> str:
    """
    加载单个文档并返回纯文本内容。

    Args:
        file_path: 文档路径

    Returns:
        文档的纯文本内容
    """
    file_path = Path(file_path)
    ext = file_path.suffix.lower()

    if ext == ".pdf":
        return read_pdf(file_path)
    elif ext in (".md", ".markdown"):
        return read_markdown(file_path)
    elif ext in (".html", ".htm"):
        return read_html(file_path)
    elif ext == ".txt":
        return read_txt(file_path)
    elif ext == ".docx":
        return read_docx(file_path)
    else:
        raise ValueError(f"不支持的文件格式: {ext}")


def iter_documents(directory: str | Path) -> Iterator[tuple[Path, str]]:
    """
    遍历目录下所有支持的文档，逐个加载。

    Args:
        directory: 文档目录

    Yields:
        (文件路径, 纯文本内容)
    """
    directory = Path(directory)
    if not directory.exists():
        raise FileNotFoundError(f"目录不存在: {directory}")

    for file_path in sorted(directory.rglob("*")):
        if file_path.is_file() and file_path.suffix.lower() in SUPPORTED_EXTENSIONS:
            try:
                text = load_document(file_path)
                if text.strip():
                    yield file_path, text
                    logger.info(f"✅ 已加载: {file_path.name} ({len(text)} 字符)")
                else:
                    logger.warning(f"⚠️  文件为空，跳过: {file_path.name}")
            except Exception as e:
                logger.error(f"❌ 加载失败 {file_path.name}: {e}")


# ---------------------------------------------------------------------------
# LlamaIndex 集成
# ---------------------------------------------------------------------------

def load_documents_to_chroma(
    input_dir: str | Path,
    persist_dir: str | Path | None = None,
) -> int:
    """
    从目录加载所有文档，分块、向量化并写入 Chroma。

    Args:
        input_dir: 文档输入目录
        persist_dir: Chroma 持久化目录（默认使用 config 中的值）

    Returns:
        导入的文本块数量
    """
    from llama_index.core import StorageContext, VectorStoreIndex, Document as LlamaDocument
    from llama_index.vector_stores.chroma import ChromaVectorStore
    from llama_index.core.node_parser import SentenceSplitter
    import chromadb

    from .config import get_config, init_settings
    init_settings()  # 确保 LlamaIndex 全局 Settings 已配置好 embed_model
    cfg = get_config()

    if persist_dir is None:
        persist_dir = cfg.absolute_chroma_dir
    persist_dir = str(persist_dir)

    logger.info(f"📚 正在扫描文档目录: {input_dir}")
    documents: list[LlamaDocument] = []

    for file_path, text in iter_documents(input_dir):
        # 用文件路径作为文档元数据
        rel_path = str(file_path.relative_to(input_dir))
        doc = LlamaDocument(
            text=text,
            metadata={
                "title": file_path.stem,
                "source": rel_path,
                "file_path": str(file_path),
            },
        )
        documents.append(doc)

    if not documents:
        logger.error("❌ 没有找到任何文档！")
        return 0

    logger.info(f"📄 共加载 {len(documents)} 份文档")

    # Chroma 持久化
    logger.info(f"💾 初始化 Chroma: {persist_dir}")
    chroma_client = chromadb.PersistentClient(path=persist_dir)
    # 如果 collection 已存在就删除（全量重建）
    try:
        chroma_client.delete_collection("linear_algebra")
    except Exception:
        pass
    chroma_collection = chroma_client.create_collection("linear_algebra")

    vector_store = ChromaVectorStore(chroma_collection=chroma_collection)
    storage_context = StorageContext.from_defaults(vector_store=vector_store)

    # 分块器
    node_parser = SentenceSplitter(
        chunk_size=cfg.chunk_size,
        chunk_overlap=cfg.chunk_overlap,
    )

    # 构建索引（会自动 embedding）
    logger.info("🔨 正在构建向量索引（本地 Embedding 模型）...")
    index = VectorStoreIndex.from_documents(
        documents,
        storage_context=storage_context,
        node_parser=node_parser,
        show_progress=True,
    )

    # 保存（LlamaIndex 默认会和 Chroma 一起持久化）
    index.storage_context.persist(persist_dir=persist_dir)

    count = chroma_collection.count()
    logger.info(f"✅ 知识库构建完成！共 {count} 个文本块")
    return count


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def cli():
    parser = argparse.ArgumentParser(description="文档加载器 - 构建知识库")
    parser.add_argument("--input", required=True, help="文档输入目录")
    parser.add_argument("--output", default=None, help="Chroma 持久化目录")
    args = parser.parse_args()

    load_documents_to_chroma(args.input, args.output)


if __name__ == "__main__":
    cli()
